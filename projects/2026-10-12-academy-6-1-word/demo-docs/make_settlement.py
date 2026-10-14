"""Mock 'their contract' for Academy 6.1: Bowater Sykes' draft settlement agreement in Nair v Calderdale.

Facts come from the mock case (names, refs, case number, Acas, dates). Written for this video only.
Run: .venv/bin/python demo-docs/make_settlement.py
"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

OUT = Path(__file__).with_name("Settlement_Agreement_Bowater_Sykes_Draft.docx")

doc = Document()
st = doc.styles["Normal"]
st.font.name = "Times New Roman"
st.font.size = Pt(12)


def para(text, bold=False, align=None, space=8):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space)
    return p


para("WITHOUT PREJUDICE AND SUBJECT TO CONTRACT", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
para("Bowater Sykes draft for discussion · Their ref: BS/CLL/2291", align=WD_ALIGN_PARAGRAPH.CENTER, space=18)
para("SETTLEMENT AGREEMENT", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space=18)

para("This agreement is made between:")
para("(1) Calderdale Logistics Limited (company number 06214477) of Brighouse Distribution Park, "
     "Clifton Common, Brighouse HD6 4HA (the Employer); and")
para("(2) Ms Priya Nair of 14 Saltaire Road, Shipley BD18 3HN (the Employee).", space=14)

clauses = [
    ("1. Background",
     "The Employee's employment ended on 31 March 2025. The Employee brought claims in the Leeds "
     "Employment Tribunal under case number 1803442/2025 (the Proceedings). The parties wish to settle "
     "all claims between them on the terms of this agreement."),
    ("2. Settlement payment",
     "The Employer shall pay the Employee the sum of £45,000 (the Settlement Payment) within 28 days of "
     "the later of the date of this agreement and receipt by the Employer of the signed adviser's "
     "certificate at clause 10."),
    ("3. Tax",
     "The Employee shall be solely responsible for any income tax or employee National Insurance "
     "contributions due on the Settlement Payment and shall indemnify the Employer against any such "
     "liability, together with any interest, penalties and costs."),
    ("4. Withdrawal of the Proceedings",
     "Within 7 days of receipt of the Settlement Payment, the Employee shall withdraw the Proceedings "
     "and consents to their dismissal by the Tribunal."),
    ("5. Confidentiality",
     "The Employee shall keep the existence and terms of this agreement, and the circumstances leading "
     "to it, strictly confidential and shall not disclose them to any person."),
    ("6. Non-disparagement",
     "The Employee shall not make, publish or communicate any statement which disparages or is "
     "derogatory of the Employer or any of its officers, employees or customers."),
    ("7. Reference",
     "On request from a prospective employer, the Employer shall provide a reference confirming only "
     "the Employee's job title and dates of employment."),
    ("8. Repayment",
     "If the Employee breaches clause 5 or clause 6, the whole of the Settlement Payment shall be "
     "repayable to the Employer on demand as a debt."),
    ("9. Full and final settlement",
     "This agreement is in full and final settlement of all claims the Employee has or may have against "
     "the Employer arising out of her employment or its termination, including the claims in the "
     "Proceedings."),
    ("10. Independent advice",
     "The Employee confirms she has received independent legal advice on the terms and effect of this "
     "agreement from Adam Frith of Aldridge Vine LLP, a relevant independent adviser, who will sign the "
     "certificate in the schedule."),
    ("11. Governing law",
     "This agreement is governed by the law of England and Wales."),
]
for head, body in clauses:
    para(head, bold=True, space=4)
    para(body, space=12)

para("Signed for and on behalf of Calderdale Logistics Limited ............................", space=10)
para("Signed by Ms Priya Nair ............................", space=10)

doc.save(OUT)
print(OUT)
